"""Local engine benchmark. Personal inputs never leave this machine.

Run setup, pdf2docx, then inspect after the JS/Word/LibreOffice stages.
Development dependencies: the existing conversion-venv (not app dependencies).
"""
import collections
import hashlib
import html
import importlib.metadata
import json
import re
import shutil
import sys
import time
import unicodedata
from pathlib import Path
from zipfile import ZipFile
from xml.etree import ElementTree as ET

import fitz
import numpy as np
from docx import Document
from docx.shared import Inches, Pt, RGBColor

ROOT = Path('.tools/professional-conversion')
W = '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'


def digest(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def save_json(path, value):
    path.write_text(json.dumps(value, indent=2, ensure_ascii=True), encoding='utf-8')


def setup():
    sources = ROOT / 'sources'
    sources.mkdir(parents=True, exist_ok=True)
    originals = []
    for name in ['Muhammad_Dawood (1).pdf', 'Muhammad_Dawood.docx']:
        candidates = [Path('tests/fixtures') / name, Path('test/fileflow_conversion_test_files') / name]
        path = next(p for p in candidates if p.exists())
        originals.append({'path': str(path), 'sha256': digest(path)})
        shutil.copyfile(path, sources / ('primary-cv.pdf' if name.endswith('.pdf') else 'supplied-cv.docx'))
    for source in Path('.tools/pdf-word-quality/sources').glob('*.pdf'):
        if source.name != 'primary-cv.pdf':
            shutil.copyfile(source, sources / source.name)

    # A ruled grid with wrapped text and a horizontally merged heading cell.
    d = fitz.open(); p = d.new_page(width=612, height=792)
    p.insert_text((50, 60), 'MERGED AND MULTILINE TABLE', fontsize=18, fontname='hebo')
    for y in [100, 135, 195, 230]:
        p.draw_line((50, y), (560, y), width=.75)
    for x in [50, 560]:
        p.draw_line((x, 100), (x, 230), width=.75)
    for x in [265, 410]:
        p.draw_line((x, 135), (x, 230), width=.75)
    for x, y, value in [(58, 122, 'Quarterly summary'), (58, 155, '001234'),
                        (273, 155, '42.50'), (418, 155, 'First cell line'),
                        (418, 170, 'Second cell line'), (58, 217, '009876'),
                        (273, 217, '87.00'), (418, 217, 'Another row')]:
        p.insert_text((x, y), value, fontsize=10)
    p.insert_text((50, 280), 'Text below the table remains independently editable.', fontsize=11)
    d.save(sources / 'merged-table.pdf'); d.close()

    # Authored DOCX: tests genuine office layout, not a PDF reconstruction.
    doc = Document(); sec = doc.sections[0]
    sec.page_width = Inches(8.5); sec.page_height = Inches(11)
    sec.top_margin = sec.bottom_margin = Inches(.7)
    sec.left_margin = sec.right_margin = Inches(.8)
    doc.styles['Normal'].font.name = 'Arial'; doc.styles['Normal'].font.size = Pt(11)
    sec.header.paragraphs[0].text = 'FileFlow synthetic office rendering reference'
    sec.footer.paragraphs[0].text = 'Offline synthetic fixture - no personal information'
    doc.add_heading('Structured office report', 0)
    p = doc.add_paragraph('A styled paragraph with ')
    p.add_run('bold').bold = True; p.add_run(', ')
    r = p.add_run('italic blue text'); r.italic = True; r.font.color.rgb = RGBColor(0x12, 0x45, 0x89)
    p.add_run(' and wrapping that follows the authored page margins. ' * 3)
    table = doc.add_table(rows=4, cols=3); table.style = 'Table Grid'
    table.cell(0, 0).merge(table.cell(0, 2)).text = 'Merged quarterly summary'
    for row, values in zip(table.rows[1:], [('ID', 'Amount', 'Notes'), ('001234', '42.50', 'A multiline cell\nSecond editable line'), ('009876', '87.00', 'Another row')]):
        for cell, value in zip(row.cells, values): cell.text = value
    # Independent image, never a whole-page facsimile.
    pix = fitz.Pixmap(fitz.csRGB, 80, 40, bytes([32, 95, 155]) * 80 * 40, False)
    image = sources / 'office-image.png'; pix.save(image)
    doc.add_picture(str(image), width=Inches(1.2))
    doc.add_page_break(); doc.add_heading('Explicit second page', 1)
    doc.add_paragraph('Page two must follow the explicit Word page break.')
    doc.save(sources / 'office-report.docx')
    save_json(ROOT / 'originals.json', originals)
    snapshot = Path('web/.pdf-word-benchmark/professional')
    snapshot.mkdir(parents=True, exist_ok=True)
    for name in ['localDocuments.ts', 'pdfWordLayout.ts', 'pdfWordExport.ts', 'pdfWordPaint.ts', 'localWordPdf.ts']:
        if not (snapshot / name).exists():
            shutil.copyfile(Path('web/src/lib') / name, snapshot / name)
    save_json(ROOT / 'baseline-snapshot.json', {p.name: digest(p) for p in snapshot.glob('*.ts')})
    print('Prepared local sources and a complete pre-change browser engine snapshot.')


def pdf2docx():
    from pdf2docx import Converter
    output = ROOT / ('pdf2docx-latest' if sys.argv[1] == 'pdf2docx-latest' else 'pdf2docx'); output.mkdir(exist_ok=True)
    records = []
    for source in sorted((ROOT / 'sources').glob('*.pdf')):
        if len(sys.argv) > 2 and source.name != sys.argv[2]:
            continue
        start = time.perf_counter()
        converter = Converter(str(source))
        try:
            converter.convert(str(output / (source.stem + '.docx')))
            records.append({'source': source.name, 'sha256': digest(source), 'seconds': time.perf_counter() - start, 'status': 'generated'})
        except Exception as error:
            records.append({'source': source.name, 'status': 'failed', 'error': str(error)})
        finally:
            converter.close()
    save_json(output / 'generation.json', {'version': importlib.metadata.version('pdf2docx'), 'PyMuPDF': importlib.metadata.version('PyMuPDF'), 'results': records})


def normalize(text):
    return re.sub(r'\s+', '', unicodedata.normalize('NFKC', text))


def render_records(value):
    # PowerShell 5 can wrap a pipeline array as {value: [...], Count: n}.
    # Recover the original records; their DOCX/PDF checksums are still verified.
    if isinstance(value, list):
        return [record for child in value for record in render_records(child)]
    if isinstance(value, dict) and 'name' in value:
        return [value]
    if isinstance(value, dict) and 'value' in value:
        return render_records(value['value'])
    raise ValueError('Invalid Word render manifest')


def structure(path):
    with ZipFile(path) as z:
        parts = [ET.fromstring(z.read(name)) for name in z.namelist() if re.fullmatch(r'word/(?:document|header\d*|footer\d*)\.xml', name)]
        mc = '{http://schemas.openxmlformats.org/markup-compatibility/2006}'
        for part in parts:
            for alternate in part.iter(mc + 'AlternateContent'):
                # Word renders a supported Choice, not its duplicate VML Fallback.
                if alternate.find(mc + 'Choice') is not None:
                    for fallback in alternate.findall(mc + 'Fallback'):
                        alternate.remove(fallback)
        def paragraph_text(paragraph):
            texts = []
            def walk(node):
                for child in node:
                    if child.tag == W + 'p': continue  # A nested text box has its own paragraphs.
                    if child.tag == W + 't': texts.append(child.text or '')
                    else: walk(child)
            walk(paragraph)
            return ''.join(texts)
        paragraphs = [p for part in parts for p in part.iter(W + 'p') if paragraph_text(p)]
        text = ''.join(t.text or '' for part in parts for t in part.iter(W + 't'))
        return {'text_paragraphs': len(paragraphs), 'max_paragraph_characters': max((len(paragraph_text(p)) for p in paragraphs), default=0),
                'tables': sum(len(list(p.iter(W + 'tbl'))) for p in parts),
                'grid_spans': sum(len(list(p.iter(W + 'gridSpan'))) for p in parts),
                'horizontally_merged_cells': sum(int(n.get(W + 'val', '1')) > 1 for p in parts for n in p.iter(W + 'gridSpan')),
                'vertical_merge_starts': sum(n.get(W + 'val') == 'restart' for p in parts for n in p.iter(W + 'vMerge')),
                'frames': sum(len(list(p.iter(W + 'framePr'))) for p in parts),
                'textboxes': sum(len(list(p.iter(W + 'txbxContent'))) for p in parts),
                'images': len([n for n in z.namelist() if n.startswith('word/media/') and not n.endswith('/')]),
                'fonts': sorted({n.get(W + 'ascii', '') for p in parts for n in p.iter(W + 'rFonts')}),
                'font_sizes_halfpoints': sorted({int(n.get(W + 'val')) for p in parts for n in p.iter(W + 'sz') if n.get(W + 'val', '').isdigit()}),
                'editable_text': text}


def compare(source, output, directory):
    directory.mkdir(parents=True, exist_ok=True)
    atext, btext = normalize(''.join(p.get_text() for p in source)), normalize(''.join(p.get_text() for p in output))
    a, b = collections.Counter(atext), collections.Counter(btext)
    result = {'source_pages': len(source), 'output_pages': len(output), 'normalized_sequence_equal': atext == btext,
              'missing_characters': sum((a-b).values()), 'extra_characters': sum((b-a).values()), 'pages': []}
    def styles(document):
        spans = [s for p in document for b in p.get_text('dict')['blocks'] if b['type'] == 0 for line in b['lines'] for s in line['spans']]
        return {'fonts': sorted({s['font'] for s in spans}), 'sizes_pt': sorted({round(s['size'], 3) for s in spans}), 'colors_rgb': sorted({s['color'] for s in spans})}
    result['source_typography'] = styles(source)
    result['rendered_typography'] = styles(output)
    for i, page in enumerate(source):
        pix = page.get_pixmap(matrix=fitz.Matrix(1.5, 1.5), alpha=False)
        pix.save(directory / f'{i+1}-source.png')
        if i >= len(output): continue
        other = output[i].get_pixmap(matrix=fitz.Matrix(1.5, 1.5), alpha=False)
        other.save(directory / f'{i+1}-output.png')
        if (pix.width, pix.height) != (other.width, other.height):
            result['pages'].append({'page': i+1, 'dimension_mismatch': True}); continue
        x = np.frombuffer(pix.samples, np.uint8).reshape(pix.height, pix.width, 3)
        y = np.frombuffer(other.samples, np.uint8).reshape(other.height, other.width, 3)
        diff = np.abs(x.astype(np.int16) - y.astype(np.int16))
        ma, mb = x.mean(axis=2) < 220, y.mean(axis=2) < 220
        union = (ma | mb).sum()
        result['pages'].append({'page': i+1, 'rgb_error': round(float(diff.mean()), 4), 'foreground_iou': round(float((ma & mb).sum()/union), 4) if union else 1})
        # Contact sheet shows actual whole pages with no registration or rescaling.
        sheet = np.concatenate([x, y], axis=1)
        fitz.Pixmap(fitz.csRGB, sheet.shape[1], sheet.shape[0], sheet.tobytes(), False).save(directory / f'{i+1}-side-by-side.png')
    return result


def inspect():
    originals = json.loads((ROOT / 'originals.json').read_text())
    assert all(digest(Path(o['path'])) == o['sha256'] for o in originals), 'An original fixture changed!'
    report = {'originals_unchanged': True, 'scope': 'Local samples only. RGB error is not a general accuracy percentage.',
              'source_hashes': {p.name: digest(p) for p in (ROOT / 'sources').iterdir() if p.suffix in ('.pdf', '.docx')}, 'samples': {}, 'word_to_pdf': {}}
    cards = []
    for source in sorted((ROOT / 'sources').glob('*.pdf')):
        engines = {}
        for engine in ['browser-before', 'browser-after', 'pdf2docx', 'pdf2docx-latest', 'word-import', 'libreoffice-import', 'native-writer']:
            docx = ROOT / engine / (source.stem + '.docx')
            rendered = docx.with_suffix('.pdf')
            if not docx.exists(): continue
            metrics = structure(docx)
            if rendered.exists():
                with fitz.open(source) as a, fitz.open(rendered) as b:
                    metrics.update(compare(a, b, ROOT / 'images' / source.stem / engine))
                records = ROOT / engine / 'word-render.json'
                if records.exists():
                    entries = render_records(json.loads(records.read_text(encoding='utf-8-sig')))
                    record = next((r for r in entries if r['name'] == docx.name), {})
                    metrics['render_matches_docx_hash'] = record.get('sourceSha256') == digest(docx) and record.get('pdfSha256') == digest(rendered)
                cards.append(f'<article><h2>{html.escape(source.stem)} / {engine}</h2><p>Original PDF on left; Word rendering on right.</p><img src="images/{source.stem}/{engine}/1-side-by-side.png"></article>')
            engines[engine] = metrics
        report['samples'][source.stem] = engines
    for name in ['supplied-cv', 'office-report']:
        reference = ROOT / 'word-reference' / (name + '.pdf')
        if not reference.exists(): continue
        manifest = render_records(json.loads((ROOT / 'word-reference/word-render.json').read_text(encoding='utf-8-sig')))
        reference_record = next(r for r in manifest if r['name'] == name + '.docx')
        assert reference_record['sourceSha256'] == digest(ROOT / 'sources' / (name + '.docx')), 'Stale Word reference input'
        assert reference_record['pdfSha256'] == digest(reference), 'Stale Word reference PDF'
        engines = {}
        for engine in ['word-pdf-before', 'word-pdf-after', 'libreoffice', 'office-button', 'pymupdf-office']:
            output = ROOT / engine / (name + '.pdf')
            if not output.exists(): continue
            with fitz.open(reference) as a, fitz.open(output) as b:
                engines[engine] = compare(a, b, ROOT / 'images' / name / engine)
            cards.append(f'<article><h2>{name} / {engine}</h2><p>Microsoft Word reference on left; converter PDF on right.</p><img src="images/{name}/{engine}/1-side-by-side.png"></article>')
        report['word_to_pdf'][name] = engines
    with fitz.open(ROOT / 'sources/primary-cv.pdf') as source, fitz.open(ROOT / 'word-reference/supplied-cv.pdf') as historical, fitz.open(ROOT / 'word-pdf-before/supplied-cv.pdf') as basic:
        supplied = structure(ROOT / 'sources/supplied-cv.docx')
        supplied.update(compare(source, historical, ROOT / 'images/primary-cv/supplied-historical'))
        report['historical_cv'] = supplied
        provenance = {'source_sha256': digest(ROOT / 'sources/primary-cv.pdf'), 'supplied_docx_sha256': digest(ROOT / 'sources/supplied-cv.docx'),
                      'source_pdf_metadata': source.metadata, 'basic_word_pdf_matches_supplied_pdf': compare(source, basic, ROOT / 'images/primary-cv/basic-roundtrip'),
                      'interpretation': 'The supplied PDF has the same rendered appearance as basic Word-to-PDF reflow of the supplied DOCX. This supports a reflowed-derivative inference; metadata does not establish which application or workflow created the file.'}
        report['fixture_provenance'] = provenance
        save_json(ROOT / 'fixture-provenance.json', provenance)
        cards.insert(0, '<article><h2>Supplied historical CV DOCX</h2><p>Supplied PDF on left; fresh Word rendering of the supplied previous DOCX on right. The current FileFlow CV result predates this task and is unchanged by its table repair.</p><img src="images/primary-cv/supplied-historical/1-side-by-side.png"></article>')
    save_json(ROOT / 'comparison.json', report)
    for name in ['merged-table', 'office-export']:
        if name not in report['samples']:
            continue
        # Presentation crop only; measurements above always use the full original pages.
        panels = [('Original PDF', ROOT / 'images' / name / 'browser-after/1-source.png'),
                  ('FileFlow at task start', ROOT / 'images' / name / 'browser-before/1-output.png'),
                  ('FileFlow after changes', ROOT / 'images' / name / 'browser-after/1-output.png')]
        figure = fitz.open(); canvas = figure.new_page(width=612*3, height=354)
        for index, (label, path) in enumerate(panels):
            if not path.exists(): continue
            pix = fitz.Pixmap(path)
            data = np.frombuffer(pix.samples, np.uint8).reshape(pix.height, pix.width, pix.n)[:495, :, :3].copy()
            crop = fitz.Pixmap(fitz.csRGB, data.shape[1], data.shape[0], data.tobytes(), False)
            canvas.insert_text((index*612+15, 17), label, fontsize=12)
            canvas.insert_image(fitz.Rect(index*612, 24, (index+1)*612, 354), stream=crop.tobytes('png'))
        canvas.get_pixmap(matrix=fitz.Matrix(1.5, 1.5), alpha=False).save(ROOT / (name + '-before-after.png'))
        figure.close()
    (ROOT / 'comparison.html').write_text('<!doctype html><meta charset="utf-8"><title>Local conversion engine comparison</title><style>body{font:16px system-ui;background:#eef1f5;padding:24px;margin:auto;max-width:1600px}article{background:white;padding:20px;margin:24px 0}img{width:100%}</style><h1>FileFlow local engine comparison</h1><p>Personal fixtures remain local. Images use identical 108 dpi, without alignment correction. Editing and rendering are different tests.</p>' + ''.join(cards), encoding='utf-8')
    if '--assert' in sys.argv:
        checks = []
        for name, engines in report['samples'].items():
            before, after = engines['browser-before'], engines['browser-after']
            assert before.get('render_matches_docx_hash') and after.get('render_matches_docx_hash'), f'{name}: stale or missing Word render'
            assert after['source_pages'] == after['output_pages'], f'{name}: pagination changed'
            assert after['missing_characters'] == after['extra_characters'] == 0, f'{name}: text lost or duplicated'
            assert sum(p['rgb_error'] for p in after['pages']) <= sum(p['rgb_error'] for p in before['pages']) + .0001, f'{name}: image error increased'
            assert sum(p['foreground_iou'] for p in after['pages']) >= sum(p['foreground_iou'] for p in before['pages']) - .0001, f'{name}: foreground overlap decreased'
            checks.append(name + ': current source/output hashes, same page count, no rendered text loss, no visual regression')
        merged = report['samples']['merged-table']
        assert merged['browser-before']['tables'] == 0
        assert merged['browser-after']['tables'] == 1 and merged['browser-after']['horizontally_merged_cells'] == 1
        assert merged['browser-after']['pages'][0]['rgb_error'] < merged['browser-before']['pages'][0]['rgb_error']
        edited = json.loads((ROOT / 'editability/verification.json').read_text(encoding='utf-8-sig'))
        assert edited['result'] == 'PASS' and edited['editedTextFound'] and edited['cells'] == 7
        native = merged['native-writer']
        assert native['render_matches_docx_hash'] and native['tables'] == 1 and native['horizontally_merged_cells'] == 1
        assert native['missing_characters'] == native['extra_characters'] == 0 and native['output_pages'] == 1
        service = json.loads((ROOT / 'office-button/verification.json').read_text())
        assert service['result'] == 'PASS' and service['failedServiceHasNoTextFallback']
        buttons = json.loads((ROOT / 'pdf-word-button/verification.json').read_text())
        assert buttons['result'] == 'PASS' and len(buttons['records']) == 3
        for name, engines in report['word_to_pdf'].items():
            actual = engines['office-button']
            assert actual['source_pages'] == actual['output_pages'] and actual['missing_characters'] == actual['extra_characters'] == 0, name + ': office renderer lost content'
        save_json(ROOT / 'quality-gates.json', {'result': 'PASS', 'checks': checks + ['Word table-cell edit saved and rendered', 'Native compiled writer table rendered; Android extraction unverified', 'React PDF-to-Word buttons tested on CV and both merged tables, no uploads', 'React button explicitly opts into real loopback office rendering; no silent text fallback'], 'originals_unchanged': True})
    print(json.dumps({name: {engine: {k: m.get(k) for k in ['text_paragraphs', 'tables', 'grid_spans', 'source_pages', 'output_pages', 'missing_characters', 'extra_characters', 'pages']} for engine, m in engines.items()} for name, engines in report['samples'].items()}, indent=2))


if __name__ == '__main__':
    {'setup': setup, 'pdf2docx': pdf2docx, 'pdf2docx-latest': pdf2docx, 'inspect': inspect}[sys.argv[1]]()

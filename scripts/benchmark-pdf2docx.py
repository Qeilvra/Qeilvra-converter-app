"""Synthetic structural benchmark for the existing engine, not a visual fidelity claim.

Run with the Python environment containing server/requirements.txt.
All samples and outputs live in a temporary directory and are deleted afterward.
"""
import contextlib
import importlib.metadata
import io
import json
import subprocess
import sys
import tempfile
import time
from pathlib import Path

import fitz
from docx import Document
from pdf2docx import Converter


def editable_text(document):
    texts = [p.text for p in document.paragraphs]
    for table in document.tables:
        texts.extend(cell.text for row in table.rows for cell in row.cells)
    return "\n".join(texts)


def create_fixture(path, scenario):
    pdf = fitz.open()
    page = pdf.new_page(width=595, height=842)
    expected = []

    def text(point, value, font="helv", size=12):
        page.insert_text(point, value, fontsize=size, fontname=font)
        expected.append(value)

    text((45, 60), "FileFlow structural benchmark", "hebo", 20)
    if scenario == "styled-multipage":
        text((45, 100), "Bold heading", "hebo", 16)
        text((45, 130), "Italic paragraph with editable content.", "heit", 12)
        text((45, 160), "Normal paragraph with 001234 and 12345678901234567890.")
        text((45, 805), "Footer page one", size=9)
        page = pdf.new_page(width=595, height=842)
        text((45, 60), "Second page heading", "hebo", 16)
        text((45, 100), "Page two searchable text.")
    elif scenario == "table-image":
        rows = [["ID", "Amount"], ["001234", "92.50"], ["12345678901234567890", "87.00"]]
        for y in [130, 165, 200, 235]:
            page.draw_line((45, y), (550, y))
        for x in [45, 345, 550]:
            page.draw_line((x, 130), (x, 235))
        for r, row in enumerate(rows):
            for c, value in enumerate(row):
                text((55 + 300 * c, 153 + 35 * r), value)
        # Real embedded raster, independent of any image generation service.
        image_pdf = fitz.open()
        image_page = image_pdf.new_page(width=100, height=60)
        image_page.draw_rect(fitz.Rect(5, 5, 95, 55), color=(0, 0, 1), fill=(0.1, 0.4, 0.9))
        page.insert_image(fitz.Rect(45, 280, 245, 400), stream=image_page.get_pixmap().tobytes("png"))
        image_pdf.close()
    elif scenario == "two-columns":
        for index in range(5):
            text((45, 120 + index * 25), f"Left column line {index + 1}")
            text((320, 120 + index * 25), f"Right column line {index + 1}")
    pdf.save(path)
    pages = len(pdf)
    pdf.close()
    return expected, pages


results = []
with tempfile.TemporaryDirectory(prefix="fileflow-quality-") as directory:
    for scenario in ["styled-multipage", "table-image", "two-columns"]:
        source = Path(directory) / f"{scenario}.pdf"
        output = Path(directory) / f"{scenario}.docx"
        expected, pages = create_fixture(source, scenario)
        start = time.perf_counter()
        if "--cli" in sys.argv:
            subprocess.run([sys.executable, "-m", "pdf2docx.main", "convert", str(source), str(output)],
                           check=True, capture_output=True, timeout=120)
        else:
            converter = Converter(str(source))
            try:
                with contextlib.redirect_stdout(io.StringIO()):
                    converter.convert(str(output))
            finally:
                converter.close()
        document = Document(output)
        recovered = editable_text(document)
        missing = [value for value in expected if value not in recovered]
        results.append({
            "scenario": scenario,
            "source_pages": pages,
            "input_bytes": source.stat().st_size,
            "output_bytes": output.stat().st_size,
            "elapsed_ms": round((time.perf_counter() - start) * 1000),
            "expected_text_fragments": len(expected),
            "recovered_text_fragments": len(expected) - len(missing),
            "missing_fragments": missing,
            "tables": len(document.tables),
            "inline_images": len(document.inline_shapes),
            "sections": len(document.sections),
            "semantic_headers": sum(bool(s.header.paragraphs[0].text) for s in document.sections),
            "semantic_footers": sum(bool(s.footer.paragraphs[0].text) for s in document.sections),
            "bold_runs": sum(bool(run.bold) for p in document.paragraphs for run in p.runs),
            "italic_runs": sum(bool(run.italic) for p in document.paragraphs for run in p.runs),
        })
print(json.dumps({
    "versions": {package: importlib.metadata.version(package) for package in ["pdf2docx", "PyMuPDF", "python-docx"]},
    "invocation": "python -m pdf2docx.main convert" if "--cli" in sys.argv else "Converter API",
    "scope": "Synthetic editable-text and package structure only; no rendered Word comparison or Urdu/RTL certification.",
    "results": results,
}, indent=2))

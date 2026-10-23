"""Author synthetic DOCX rendering fixtures; never reads or changes private documents."""
import json
from pathlib import Path
from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
import fitz

ROOT = Path('.tools/word-pdf-rendering')
SOURCES = ROOT / 'sources'
SOURCES.mkdir(parents=True, exist_ok=True)
logo = fitz.open(); page = logo.new_page(width=240, height=90)
page.draw_rect(page.rect, color=None, fill=(.086, .23, .38))
page.insert_text((20, 50), 'FILEFLOW / SYNTHETIC', fontsize=15, color=(1,1,1))
page.get_pixmap().save(SOURCES / 'logo.png'); logo.close()

def new(title):
    d = Document()
    s = d.sections[0]
    s.page_width, s.page_height = Inches(8.5), Inches(11)
    s.top_margin = s.bottom_margin = Inches(.65)
    s.left_margin = s.right_margin = Inches(.8)
    d.styles['Normal'].font.name = 'Arial'
    d.styles['Normal'].font.size = Pt(11)
    d.core_properties.title = title
    d.core_properties.author = 'FileFlow synthetic benchmark'
    return d

d = new('Professional CV')
p = d.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run('ALEX MORGAN'); r.bold = True; r.font.name = 'Garamond'; r.font.size = Pt(28)
p = d.add_paragraph('DOCUMENT ENGINEER | SYNTHETIC CV'); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
for heading, text in [('PROFILE', 'Document engineer experienced in typography, page layout and accessible documents. ' * 4),
                      ('EXPERIENCE', 'Rendering Engineer | Example Labs | 2022 - 2026'),
                      ('SELECTED WORK', 'Built a multilingual publishing workflow with careful font metrics and regression checks. ' * 5),
                      ('EDUCATION', 'BSc Computer Science | Example University'),
                      ('SKILLS', 'Java, document layout, tables, Unicode, PDF, typesetting')]:
    p = d.add_paragraph(); p.paragraph_format.space_before = Pt(16)
    r = p.add_run(heading); r.bold = True; r.font.size = Pt(13); r.font.color.rgb = RGBColor.from_string('163B60')
    p = d.add_paragraph(text); p.paragraph_format.line_spacing = 1.2
d.save(SOURCES / 'professional-cv.docx')

d = new('Columns and sections')
d.add_heading('Two-column technical newsletter', 0)
s = d.add_section(WD_SECTION_START.CONTINUOUS)
cols = s._sectPr.find(qn('w:cols')); cols.set(qn('w:num'), '2'); cols.set(qn('w:space'), '360')
for i in range(10):
    d.add_heading(f'Section {i+1}', 2)
    p = d.add_paragraph(('Column flow uses original margins and paragraph metrics. ' * 7))
    p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    for r in p.runs: r.font.name = 'Times New Roman'
d.save(SOURCES / 'columns.docx')

d = new('Complex office report')
d.sections[0].header.paragraphs[0].text = 'FILEFLOW | Engineering report'
d.sections[0].footer.paragraphs[0].text = 'Synthetic fixture - document footer'
d.add_picture(str(SOURCES / 'logo.png'), width=Inches(1.8))
d.add_heading('Quarterly engineering report', 0)
t = d.add_table(rows=5, cols=3); t.style = 'Table Grid'
t.cell(0,0).merge(t.cell(0,2)).text = 'Merged summary: Q4'
t.cell(1,0).merge(t.cell(2,0)).text = 'Vertical merge'
for i in range(1,5):
    t.cell(i,1).text = f'{i * 1250}.00'
    t.cell(i,2).text = 'Multiple lines\nPreserved inside cell'
for cell in t.rows[0].cells:
    shade = OxmlElement('w:shd'); shade.set(qn('w:fill'), 'DCEAF7'); cell._tc.get_or_add_tcPr().append(shade)
for i in range(4): d.add_paragraph(f'Engineering item {i+1}', style='List Number')
p = d.add_paragraph('Documentation: ')
link = OxmlElement('w:hyperlink'); link.set(qn('r:id'), d.part.relate_to('https://example.org/fileflow', 'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink', is_external=True))
run = OxmlElement('w:r'); text = OxmlElement('w:t'); text.text = 'Reference link'; run.append(text); link.append(run); p._p.append(link)
d.add_page_break(); d.add_heading('Explicit second page', 1)
d.add_paragraph('This section must start after a page break. The header and footer must remain visible.')
d.save(SOURCES / 'complex-report.docx')

d = new('Urdu and bidirectional text')
d.add_heading('Unicode and Urdu', 0)
for text in ['یہ ایک آزمائشی دستاویز ہے۔', 'اردو متن اور انگریزی الفاظ FileFlow 2026 ایک ساتھ۔', 'درست حروف، ترتیب اور دائیں سے بائیں صف بندی ضروری ہے۔']:
    p = d.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    bidi = OxmlElement('w:bidi'); p._p.get_or_add_pPr().append(bidi)
    r = p.add_run(text); r.font.size = Pt(16); r.font.name = 'Arial'
    fonts = r._r.get_or_add_rPr().find(qn('w:rFonts')); fonts.set(qn('w:cs'), 'Arial')
    rtl = OxmlElement('w:rtl'); r._r.get_or_add_rPr().append(rtl)
d.add_paragraph('Latin comparison: café — résumé — naïve. العربية. Ελληνικά.')
d.save(SOURCES / 'urdu-rtl.docx')

d = new('Shapes and text boxes')
d.add_heading('Positioned shapes and text boxes', 0)
d.add_paragraph('Body text must coexist with the positioned callout below.')
p = d.add_paragraph()
p._p.append(parse_xml('''<w:r xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" xmlns:v="urn:schemas-microsoft-com:vml"><w:pict><v:rect style="position:absolute;margin-left:110pt;margin-top:40pt;width:240pt;height:100pt" fillcolor="#e0edf9" strokecolor="#163b60"><v:textbox inset="8pt,8pt,8pt,8pt"><w:txbxContent><w:p><w:r><w:t>EDITABLE POSITIONED CALLOUT</w:t></w:r></w:p><w:p><w:r><w:t>Text inside a coloured shape.</w:t></w:r></w:p></w:txbxContent></v:textbox></v:rect></w:pict></w:r>'''))
d.add_paragraph('Following paragraph. ' * 35)
d.save(SOURCES / 'shapes.docx')

(ROOT / 'fixtures.json').write_text(json.dumps({'inputs': [p.name for p in sorted(SOURCES.glob('*.docx'))], 'origin': 'Authored synthetic DOCX, not reconstructed from PDF'}, indent=2))
print('Created', len(list(SOURCES.glob('*.docx'))), 'synthetic DOCX inputs')

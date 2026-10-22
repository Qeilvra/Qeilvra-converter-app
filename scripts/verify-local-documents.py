"""Independently inspect actual browser-generated files. Uses the optional benchmark venv."""
import json
import sys
from pathlib import Path
from docx import Document
import fitz

directory = Path(sys.argv[1] if len(sys.argv) > 1 else '.tools/backend-free-evidence/chromium-desktop')
word = Document(directory / 'editable.docx')
pdf = fitz.open(directory / 'source.pdf')
assert len(word.tables) == 2
assert len(word.sections) == 2
assert len(word.inline_shapes) == 2
for table in word.tables:
    assert table.cell(1, 0).text == '001234'
    assert table.cell(2, 0).text == '12345678901234567890'
    assert table.cell(1, 1).text == '$1,234.50'
    assert table.cell(2, 1).text == '=SUM(A1:A2)'
assert sum('This sentence must remain editable' in p.text for p in word.paragraphs) == 2
for section, page in zip(word.sections, pdf):
    assert abs(section.page_width.pt - page.rect.width) < 1
    assert abs(section.page_height.pt - page.rect.height) < 1
assert any(run.bold and run.font.size and run.font.size.pt == 18 for p in word.paragraphs for run in p.runs), [(p.text, [(r.font.name, r.font.size.pt if r.font.size else None, r.bold) for r in p.runs]) for p in word.paragraphs]
result = {'result': 'PASS', 'editable_paragraphs': len(word.paragraphs), 'tables': len(word.tables), 'bitmap_images': len(word.inline_shapes), 'sections': len(word.sections), 'source_pages': len(pdf), 'scope': 'Independent python-docx/PyMuPDF parsing; no Word rendering or pixel-layout equivalence claim'}
print(json.dumps(result, indent=2))
(directory / 'independent-inspection.json').write_text(json.dumps(result, indent=2), encoding='utf-8')
